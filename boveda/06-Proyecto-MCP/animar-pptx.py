# -*- coding: utf-8 -*-
"""Agrega transiciones, animaciones de entrada y el guion como notas a un .pptx
exportado desde la presentacion web (que las pierde al exportar).

    python3 animar.py entrada.pptx salida.pptx guion.md [tipografia]
"""
import copy, re, sys
from lxml import etree
from pptx import Presentation

ENT, SAL, GUION = sys.argv[1], sys.argv[2], sys.argv[3]
FUENTE = sys.argv[4] if len(sys.argv) > 4 else None
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
prs = Presentation(ENT)
K = prs.slide_width / 1920.0

# grupos que entran con un clic, por lamina: cajas (x1, y1, x2, y2) en px de 1920 x 1080
COLS4 = [(128, 526), (550, 948), (972, 1370), (1394, 1792)]
GRUPOS = {
    2: [(128, 460, 700, 760), (704, 460, 1276, 760), (1280, 460, 1860, 760)],
    5: [(a, 430, b, 960) for a, b in COLS4],
    6: [(128, 430, 661, 960), (693, 430, 1226, 960), (1259, 430, 1792, 960)],
    7: [(100, y, 1820, y + 116) for y in (400, 540, 680, 820)],
    8: [(a, 460, b, 960) for a, b in COLS4],
    11: [(100, 410, 1820, 830), (100, 850, 1860, 920)],
    12: [(100, y, 1820, y + 140) for y in (470, 634, 798)],
}
EMPUJE = {3, 8, 12}          # donde empieza cada bloque de la rubrica


def timing_xml(grupos):
    n = [2]

    def nid():
        n[0] += 1
        return n[0]
    clicks = []
    for ids in grupos:
        efectos = []
        for i, sp in enumerate(ids):
            # entrada "ascender": aparece con desvanecido mientras sube 6 % del alto de la lamina.
            # Todas las formas de una tarjeta llevan el mismo desplazamiento, asi que sube entera.
            dest = '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>' % sp
            efectos.append(
                '<p:par><p:cTn id="%d" presetID="42" presetClass="entr" presetSubtype="0" fill="hold" nodeType="%s">'
                '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                '<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                '%s<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                '</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="600"/>%s</p:cBhvr></p:animEffect>'
                '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="%d" dur="600" fill="hold"/>%s'
                '<p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                '<p:tav tm="0"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav>'
                '<p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim>'
                '<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="%d" dur="600" decel="100000" fill="hold"/>%s'
                '<p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                '<p:tav tm="0"><p:val><p:strVal val="#ppt_y+.06"/></p:val></p:tav>'
                '<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>'
                '</p:childTnLst></p:cTn></p:par>'
                % (nid(), "clickEffect" if i == 0 else "withEffect", nid(), dest, nid(), dest, nid(), dest, nid(), dest))
        clicks.append(
            '<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>%s'
            '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>' % (nid(), nid(), "".join(efectos)))
    return ('<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            '<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            '<p:childTnLst>%s</p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>' % (P, "".join(clicks)))


# --- el guion: el texto citado de cada seccion "## N · ..."
guion = {}
actual = None
for linea in open(GUION, encoding="utf-8").read().split("\n"):
    m = re.match(r"^## (\d+) · ", linea)
    if m:
        actual = int(m.group(1)); guion[actual] = []
    elif linea.startswith("## "):
        actual = None
    elif actual and linea.startswith(">"):
        t = linea[1:].strip().replace("*(clic)*", "[CLIC]")
        guion[actual].append(re.sub(r"\*+", "", t))
guion = {k: re.sub(r"\n{3,}", "\n\n", "\n".join(v)).strip() for k, v in guion.items()}

informe = []
for num, slide in enumerate(prs.slides, 1):
    sld = slide._element
    for viejo in sld.findall("{%s}transition" % P) + sld.findall("{%s}timing" % P):
        sld.remove(viejo)
    ext = sld.find("{%s}extLst" % P)
    nuevos = [etree.fromstring('<p:transition xmlns:p="%s" spd="med">%s</p:transition>'
                               % (P, '<p:push dir="l"/>' if num in EMPUJE else '<p:fade/>'))]
    grupos = []
    for (x1, y1, x2, y2) in GRUPOS.get(num, []):
        piezas = [sh for sh in slide.shapes
                  if x1 <= (sh.left + sh.width / 2) / K <= x2 and y1 <= (sh.top + sh.height / 2) / K <= y2]
        if len(piezas) > 1:
            # la tarjeta se agrupa en UN objeto: si cada pieza se anima por separado, Keynote las
            # hace viajar distancias distintas y se enciman a mitad de la subida
            grupo = slide.shapes.add_group_shape(shapes=piezas)
            grupo.name = "Tarjeta %d" % (len(grupos) + 1)
            grupos.append([grupo.shape_id])
        elif piezas:
            grupos.append([piezas[0].shape_id])
    if grupos:
        nuevos.append(etree.fromstring(timing_xml(grupos)))
    for el in nuevos:
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)
    if guion.get(num):
        slide.notes_slide.notes_text_frame.text = guion[num]
    informe.append((num, "push" if num in EMPUJE else "fade", len(grupos), sum(len(g) for g in grupos), bool(guion.get(num))))

if FUENTE:
    for parte in [s.part for s in prs.slides] + [l.part for l in prs.slide_layouts] + [m.part for m in prs.slide_masters]:
        for el in parte._element.iter("{%s}latin" % A, "{%s}ea" % A, "{%s}cs" % A):
            if el.get("typeface") == "DM Sans":
                el.set("typeface", FUENTE)

prs.core_properties.author = "Grupo 2"
prs.core_properties.last_modified_by = "Santiago Barrera"
prs.core_properties.title = "Caso Harley-Davidson: arquitectura del proyecto SiL'K"
prs.save(SAL)

if FUENTE:
    # el tema y los patrones tambien nombran la tipografia: se reescribe el paquete
    import os, shutil, zipfile
    tmp = SAL + ".tmp"
    with zipfile.ZipFile(SAL) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                datos = datos.replace(b'typeface="DM Sans"', ('typeface="%s"' % FUENTE).encode("utf-8"))
            zout.writestr(item, datos)
    shutil.move(tmp, SAL)
print("lámina  transición  clics  formas-animadas  guion")
for r in informe:
    print("  %2d     %-9s  %d      %2d               %s" % r)
print("clics en total:", sum(r[2] for r in informe))
